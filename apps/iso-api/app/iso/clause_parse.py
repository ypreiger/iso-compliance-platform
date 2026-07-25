"""Parse ISO documents preserving source order (PDF/DOC/DOCX)."""
from __future__ import annotations

import re
from io import BytesIO

from app.iso.parser import (
    ParsedClause,
    _clean_body,
    is_plausible_clause,
    normalize_clause_id,
    preprocess_extracted_text,
)

CLAUSE_LINE = re.compile(
    r"^(?:#{1,4}\s*)?"
    r"\.?"  # optional leading dot from RTL extractors
    r"(\d{1,2}(?:\.\d{1,2}){0,4})"
    r"(?:\s*[\.\)]?\s*)"
    r"(\S(?:.*\S)?)\s*$"
)
CLAUSE_ID_ONLY = re.compile(r"^\.?\d{1,2}(?:\.\d{1,2}){0,4}\.?$")
_DOTS_ONLY = re.compile(r"^[.\s…·‧]+$")

JUNK_LINE = re.compile(
    r"(©\s*ISO|All rights reserved|Licensed to|ANSI order|Downloaded \d/"
    r"|^\s*ISO\s*9001:\d{4}|^\s*INTERNATIONAL\s+STANDARD\s*$|^\d+\s*$"
    r"|\.{10,}|…{3,})",  # Table of contents dotted leaders
    re.IGNORECASE,
)


def clause_level(clause_id: str) -> int:
    return len(normalize_clause_id(clause_id).split("."))


def try_clause_header(line: str) -> tuple[str, str] | None:
    stripped = line.strip().lstrip(".")
    if not stripped or JUNK_LINE.search(stripped):
        return None

    # Skip table of contents lines (clause + title + dots/page number)
    # Pattern: "4.1 Understanding the organization.................5"
    if re.search(r'\.{5,}', stripped):  # Line has 5+ consecutive dots = TOC entry
        return None

    # Skip quoted references to clauses (e.g., in Annex or commentary)
    # Pattern: '4.1: "The organization shall...' or '4.1: "The organization...'
    # Handles both ASCII quotes and Unicode smart quotes
    if re.match(r'^\d+(?:\.\d+)*:\s*["\'“”‘’]', stripped):
        return None

    # TOC entries often lose dotted leaders in PDF extract but keep a trailing
    # page number: "7.1 Planning of product realization 12"
    if re.search(r"\s+\d{1,3}\s*$", stripped) and not re.search(
        r"\b(shall|must|should)\b", stripped, re.I
    ):
        return None

    match = CLAUSE_LINE.match(stripped)
    if not match:
        return None
    clause_id = normalize_clause_id(match.group(1))
    title = match.group(2).strip()

    # Clean up any remaining TOC artifacts
    title = re.sub(r'\.{3,}.*$', '', title).strip()  # Remove trailing dots and everything after
    title = re.sub(r'…{2,}.*$', '', title).strip()   # Remove ellipsis
    title = re.sub(r'\s+\d+\s*$', '', title).strip()  # Remove trailing page numbers

    # Correspondence-table contamination: "Support 6 Resource management"
    # (ISO 9001 title + ISO 13485 id/title on the same extracted line).
    if re.search(
        r"\s+\d{1,2}(?:\.\d{1,2}){0,4}\s+[A-Za-z\u0590-\u05FF]",
        title,
    ):
        return None

    # OCR / compacted PDFs glue the body onto the heading. Validate the short
    # title only; keep the full string so start_clause can peel the lead body.
    short, _lead = split_title_and_lead_body(title)

    # Body sentences mis-prefixed with a clause id must not become titles.
    if re.search(r"\bshall be\b", short, re.I) or (
        re.search(r"\b(shall|must|should)\b", short, re.I) and len(short) > 70
    ):
        return None

    if len(short) > 240:
        return None
    if not _looks_like_clause_title(short):
        return None
    if not is_plausible_clause(clause_id, short, "placeholder"):
        return None
    return clause_id, title


def split_title_and_lead_body(title: str) -> tuple[str, str]:
    t = title.strip()
    # OCR often glues a full sentence onto the heading:
    # "…before delivery The organization shall deal with…"
    glued = re.search(
        r"(?<=[A-Za-z\u0590-\u05FF\.])\s+(?="
        r"(?:The organization|When [A-Za-z]|Records of|This |Documented |"
        r"a\)\s|b\)\s|c\)\s))",
        t,
    )
    if glued and glued.start() > 4:
        head = t[: glued.start()].strip().rstrip(".")
        return head or "Requirements", t[glued.start() :].strip()
    norm = re.search(r"\b(shall|must|should|חייב|יידרש| יש )\b", t, re.IGNORECASE)
    if norm and norm.start() > 4:
        head = t[: norm.start()].strip().rstrip(".")
        tail = t[norm.start() :].strip()
        return head or "Requirements", tail
    if len(t) > 90:
        cut = t[:80].rsplit(" ", 1)[0]
        return cut, t[len(cut) :].strip()
    return t, ""


def repair_broken_words(text: str) -> str:
    stop = {
        "the", "and", "for", "with", "this", "that", "from", "are", "was", "be", "to", "of", "in",
        "on", "at", "by", "or", "an", "as", "is", "it", "its", "shall", "will", "not", "may", "can",
        "has", "have", "had", "were", "been", "being",
    }
    # Apply only to latin words. Running this on Hebrew caused valid words
    # to be merged and damaged clause readability.
    return re.sub(
        r"\b([a-z]{2,4})\s+([a-z]{4,})\b",
        lambda m: m.group(1) + m.group(2) if m.group(1) not in stop else m.group(0),
        text,
    )


_MD_HEADING = re.compile(r"^#{1,6}\s+(.*\S)\s*$")
_HEBREW_RE = re.compile(r"[\u0590-\u05FF]")


def _looks_like_clause_title(text: str) -> bool:
    """Accept short heading-like titles; reject EN body fragments used as titles."""
    t = text.strip()
    if not t or len(t) > 120:
        return False
    if t in {"סעיף", "א -", "( א -", "א-", "Introduction", "מבוא"}:
        return t in {"Introduction", "מבוא"}
    if t.endswith((";", ",")) and not _HEBREW_RE.search(t):
        return False
    # English body bullets / fragments are not titles.
    if re.match(r"^(and|or|the|of|to|for|with|from|that|which|monitoring|audit|nonconform|evaluate|records)\b", t, re.I):
        return False
    if re.match(r"^[a-z]", t):
        return False
    if re.search(r"\b(shall|must|should)\b", t, re.I):
        return False
    if re.search(r"https?://|www\.iso\.org", t, re.I):
        return False
    if not re.search(r"[A-Za-z\u0590-\u05FF]", t):
        return False
    return True


def reflow_continuation_lines(text: str) -> list[str]:
    raw_lines = [ln.strip() for ln in text.replace("\r", "\n").split("\n")]
    logical: list[str] = []
    buffer = ""
    pending_md_title = ""

    def flush_buffer() -> None:
        nonlocal buffer
        if buffer:
            logical.append(buffer)
            buffer = ""

    for line in raw_lines:
        if not line or _DOTS_ONLY.fullmatch(line):
            # Hebrew TOC dotted leaders — skip, but keep a bare clause-id buffer
            # so the next title line can still attach (5.1.2 / .... / . Title).
            if buffer and CLAUSE_ID_ONLY.fullmatch(buffer):
                continue
            flush_buffer()
            continue
        # Strip leading decorative dots from RTL PDF extracts (". Title" / ".5.1.1").
        # Do NOT strip trailing periods — those are real sentence punctuation.
        line = re.sub(r"^\.+", "", line).strip()
        if not line:
            flush_buffer()
            continue
        md = _MD_HEADING.match(line)
        if md and not try_clause_header(md.group(1)):
            pending_md_title = md.group(1).strip()
            continue
        if try_clause_header(line):
            flush_buffer()
            logical.append(line)
            pending_md_title = ""
            continue
        # Bare clause id line (common in Hebrew/RTL extracts): keep discrete so
        # the following title line can join via the id-only+title rule below.
        # Also support title-before-id (RTL): "כללי" then "5.1.1".
        if CLAUSE_ID_ONLY.fullmatch(line):
            cid = line.strip(".")
            title_candidate = ""
            if pending_md_title and _looks_like_clause_title(pending_md_title):
                title_candidate = pending_md_title
                pending_md_title = ""
            elif buffer and not try_clause_header(buffer):
                # Title-before-id is a Hebrew/RTL PDF pattern. Never steal the
                # previous English sentence as a clause title.
                if _HEBREW_RE.search(buffer):
                    parts = re.split(r"(?<=[.!?…])\s+", buffer.strip())
                    if len(parts) > 1 and _looks_like_clause_title(parts[-1]):
                        title_candidate = parts[-1].strip()
                        buffer = " ".join(parts[:-1]).strip()
                    elif _looks_like_clause_title(buffer):
                        title_candidate = buffer.strip()
                        buffer = ""
            flush_buffer()
            if title_candidate and try_clause_header(f"{cid} {title_candidate}"):
                logical.append(f"{cid} {title_candidate}")
            else:
                buffer = cid
            continue
        # Common PDF extraction pattern:
        #   4.1
        #   Understanding the organization...
        # Join these two lines into one clause header.
        if buffer and CLAUSE_ID_ONLY.fullmatch(buffer):
            if _looks_like_clause_title(line):
                joined = f"{buffer.strip('.')} {line}"
                if try_clause_header(joined):
                    logical.append(joined)
                    buffer = ""
                    continue
            # Not a title — keep id as its own header line with placeholder title
            # so subsequent body text is not lost into the previous clause.
            logical.append(f"{buffer.strip('.')} Clause {buffer.strip('.')}")
            buffer = line
            continue
        # Page chrome / copyright — never merge into the previous paragraph.
        if JUNK_LINE.search(line):
            flush_buffer()
            continue
        if not buffer:
            buffer = line
            continue
        last = buffer.split()[-1] if buffer else ""
        if last and len(last) <= 4 and last.isalpha() and line[:1].islower():
            buffer = f"{buffer}{line}"
        elif line[:1].islower() or len(line) < 48:
            buffer = f"{buffer} {line}"
        else:
            flush_buffer()
            buffer = line
    flush_buffer()
    return logical


def _append_body(store: dict[str, list[str]], clause_id: str, text: str) -> None:
    if text.strip():
        store.setdefault(clause_id, []).append(text.strip())


def parse_docx_by_structure(content: bytes) -> list[ParsedClause]:
    """Parse DOCX by heading hierarchy when clause numbers aren't in text.

    Detects heading level pattern (English uses H3/H5/H6, Hebrew uses H1/H2) and
    maps to clause numbers: top level = 1,2,3..., second level = X.1, X.2...
    """
    from docx import Document

    from app.iso.document_extract import iter_docx_blocks

    doc = Document(BytesIO(content))

    # Detect which heading levels are used by sampling
    heading_levels_found = set()
    for para in doc.paragraphs[:300]:
        if para.style and 'heading' in para.style.name.lower():
            heading_levels_found.add(para.style.name.lower())

    # Determine heading level mapping
    if 'heading 3' in heading_levels_found:
        # English pattern: H3=top, H5=sub, H6=subsub
        level_map = {'heading 3': 0, 'heading 5': 1, 'heading 6': 2}
        start_keywords = ["scope", "normative", "terms", "context"]
    else:
        # Hebrew pattern: H1=top, H2=sub
        level_map = {'heading 1': 0, 'heading 2': 1}
        start_keywords = ["חלות", "אזכור", "מונחים", "הקשר"]  # Hebrew equivalents

    clauses: list[ParsedClause] = []
    clause_counters = [0, 0, 0, 0]
    current_clause_id = ""
    body_parts: dict[str, list[str]] = {}
    seq = 0
    started = False

    for block in iter_docx_blocks(doc):
        if block.__class__.__name__ == "Paragraph":
            text = re.sub(r"\s+", " ", block.text).strip()
            if not text:
                continue

            style = (block.style.name or "").lower() if block.style else ""

            # Start numbering when we hit first main section
            if not started and style in level_map:
                text_lower = text.lower()
                if level_map[style] == 0 and any(kw in text_lower for kw in start_keywords):
                    started = True

            if not started:
                continue

            # Get heading level (0=top, 1=second, 2=third)
            if style not in level_map:
                # Regular paragraph
                if current_clause_id:
                    _append_body(body_parts, current_clause_id, text)
                continue

            level = level_map[style]

            if level == 0:
                # Top level: 1, 2, 3, 4...
                clause_counters[0] += 1
                clause_counters[1] = 0
                clause_counters[2] = 0
                current_clause_id = str(clause_counters[0])
            elif level == 1:
                # Second level: X.1, X.2...
                if clause_counters[0] == 0:
                    continue
                clause_counters[1] += 1
                clause_counters[2] = 0
                current_clause_id = f"{clause_counters[0]}.{clause_counters[1]}"
            elif level == 2:
                # Third level: X.Y.1, X.Y.2...
                if clause_counters[1] == 0:
                    continue
                clause_counters[2] += 1
                current_clause_id = f"{clause_counters[0]}.{clause_counters[1]}.{clause_counters[2]}"

            clauses.append(ParsedClause(current_clause_id, text, "", seq))
            body_parts[current_clause_id] = []
            seq += 1

        elif current_clause_id:
            # Table content
            for row in block.rows:
                cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells if c.text.strip()]
                if cells:
                    _append_body(body_parts, current_clause_id, " | ".join(cells))

    # Assign bodies
    result: list[ParsedClause] = []
    for clause in clauses:
        body = _clean_body("\n\n".join(body_parts.get(clause.clause_id, [])))
        result.append(ParsedClause(clause.clause_id, clause.title, body, clause.sort_order))

    if not result:
        raise ValueError("No clause headings found in DOCX")

    return result


def parse_docx_faithful(content: bytes) -> list[ParsedClause]:
    """Walk DOCX blocks in document order; assign paragraphs to clauses faithfully."""
    from docx import Document

    from app.iso.document_extract import iter_docx_blocks

    doc = Document(BytesIO(content))
    clauses_by_id: dict[str, ParsedClause] = {}
    order: list[str] = []
    body_parts: dict[str, list[str]] = {}
    current_id: str | None = None
    seq = 0

    def open_clause(clause_id: str, title: str) -> None:
        nonlocal current_id, seq
        short_title, lead = split_title_and_lead_body(title)
        if clause_id not in clauses_by_id:
            clauses_by_id[clause_id] = ParsedClause(clause_id, short_title, "", seq)
            order.append(clause_id)
            seq += 1
            body_parts[clause_id] = []
        else:
            clauses_by_id[clause_id].title = short_title or clauses_by_id[clause_id].title
        if lead:
            _append_body(body_parts, clause_id, lead)
        current_id = clause_id

    for block in iter_docx_blocks(doc):
        if block.__class__.__name__ == "Paragraph":
            text = re.sub(r"\s+", " ", block.text).strip()
            if not text:
                continue
            style = (block.style.name or "").lower() if block.style else ""
            header = try_clause_header(text)
            if header:
                open_clause(header[0], header[1])
                continue
            if current_id:
                _append_body(body_parts, current_id, text)
        else:
            for row in block.rows:
                cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells if c.text.strip()]
                if not cells:
                    continue
                joined = " ".join(cells)
                header = try_clause_header(joined)
                if header and len(joined) < 200:
                    open_clause(header[0], header[1])
                elif current_id:
                    _append_body(body_parts, current_id, joined)

    result: list[ParsedClause] = []
    for clause_id in order:
        clause = clauses_by_id[clause_id]
        body = repair_broken_words(_clean_body("\n\n".join(body_parts.get(clause_id, []))))
        result.append(ParsedClause(clause.clause_id, clause.title, body, clause.sort_order))
    if not result:
        raise ValueError("No clause headings found in DOC/DOCX")
    return result


def parse_hierarchical_lines(lines: list[str], *, roll_up: bool = False) -> list[ParsedClause]:
    clauses_by_id: dict[str, ParsedClause] = {}
    order: list[str] = []
    stack: list[str] = []
    body_parts: dict[str, list[str]] = {}
    seq = 0

    def start_clause(clause_id: str, title: str) -> None:
        nonlocal seq
        level = clause_level(clause_id)
        while stack and clause_level(stack[-1]) >= level:
            stack.pop()
        short_title, lead = split_title_and_lead_body(title)
        if clause_id not in clauses_by_id:
            clauses_by_id[clause_id] = ParsedClause(clause_id, short_title, "", seq)
            order.append(clause_id)
            seq += 1
            body_parts[clause_id] = []
        else:
            # First real heading wins — never let a later false header overwrite.
            if (not clauses_by_id[clause_id].title.strip()
                    or clauses_by_id[clause_id].title.startswith("Clause ")):
                if short_title:
                    clauses_by_id[clause_id].title = short_title
        if lead:
            _append_body(body_parts, clause_id, lead)
        stack.append(clause_id)

    i = 0
    while i < len(lines):
        line = lines[i]
        header = try_clause_header(line)
        if header:
            start_clause(header[0], header[1])
            i += 1
            continue
        bare = CLAUSE_ID_ONLY.fullmatch(line.strip().lstrip("."))
        if bare:
            cid = normalize_clause_id(bare.group(0))
            # OCR sometimes emits "8.3.3" then the title on the next line.
            if i + 1 < len(lines):
                nxt = lines[i + 1].strip()
                if nxt and try_clause_header(f"{cid} {nxt}"):
                    start_clause(cid, nxt)
                    i += 2
                    continue
            # Bare clause-id mid-body is usually a cross-reference / layout artifact.
            i += 1
            continue
        if stack:
            # Drop pure chrome lines; if a body line only has a trailing footer,
            # strip the footer instead of discarding the whole paragraph.
            if not (
                JUNK_LINE.fullmatch(line.strip())
                or re.match(
                    r"^(©\s*ISO|All rights reserved|Licensed to|ISO\s*\d{4,})",
                    line.strip(),
                    re.I,
                )
            ):
                cleaned = JUNK_LINE.split(line)[0].strip()
                # Drop TOC leftovers that slipped into the body stream.
                if cleaned and not (
                    re.search(r"\.{5,}", cleaned)
                    or re.match(
                        r"^\d{1,2}(?:\.\d{1,2}){0,4}\s+\S.+\s+\d{1,3}$",
                        cleaned,
                    )
                ):
                    _append_body(body_parts, stack[-1], cleaned)
        i += 1

    result: list[ParsedClause] = []
    for clause_id in order:
        clause = clauses_by_id[clause_id]
        body = repair_broken_words(_clean_body("\n\n".join(body_parts.get(clause_id, []))))
        result.append(ParsedClause(clause.clause_id, clause.title, body, clause.sort_order))

    if roll_up:
        result = roll_up_empty_parents(result)
    return result


def roll_up_empty_parents(clauses: list[ParsedClause]) -> list[ParsedClause]:
    by_id = {c.clause_id: c for c in clauses}
    children: dict[str, list[str]] = {}
    for cid in by_id:
        parts = cid.split(".")
        if len(parts) > 1:
            parent = ".".join(parts[:-1])
            if parent in by_id:
                children.setdefault(parent, []).append(cid)

    updated: list[ParsedClause] = []
    for clause in clauses:
        body = clause.body.strip()
        if not body and clause.clause_id in children:
            chunks: list[str] = []
            for child_id in sorted(children[clause.clause_id], key=lambda x: by_id[x].sort_order):
                child = by_id[child_id]
                child_body = child.body.strip() or child.title
                if child_body:
                    chunks.append(f"{child_id} {child.title}\n\n{child_body}")
            body = "\n\n".join(chunks)
        updated.append(ParsedClause(clause.clause_id, clause.title, body, clause.sort_order))
    return updated


def promote_nested_clauses(clauses: list[ParsedClause]) -> list[ParsedClause]:
    """Promote child headings that were left inside a parent body (5.1.1 inside 5.1).

    Also handles Hebrew/RTL extraction where the title line appears *before*
    a dotted clause id line:
        כללי
        .5.1.1
        body…
    """
    if not clauses:
        return clauses
    by_id: dict[str, ParsedClause] = {}
    order: list[str] = []
    bodies: dict[str, list[str]] = {}

    def ensure(cid: str, title: str, sort_order: int) -> None:
        if cid not in by_id:
            by_id[cid] = ParsedClause(cid, title, "", sort_order)
            order.append(cid)
            bodies[cid] = []
        elif title and not by_id[cid].title:
            by_id[cid].title = title

    for clause in clauses:
        ensure(clause.clause_id, clause.title, clause.sort_order)
        current = clause.clause_id
        pending_title: str | None = None
        for line in (clause.body or "").split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            id_only = CLAUSE_ID_ONLY.fullmatch(stripped.lstrip("."))
            if id_only:
                nested_id = normalize_clause_id(stripped.lstrip("."))
                if nested_id.startswith(clause.clause_id + ".") or clause_level(nested_id) > clause_level(
                    clause.clause_id
                ):
                    title = pending_title or nested_id
                    pending_title = None
                    # Previous pending title was mistakenly appended to parent — remove it.
                    if bodies.get(current) and bodies[current][-1] == title:
                        bodies[current].pop()
                    ensure(nested_id, title, _sort_key_safe(nested_id))
                    current = nested_id
                    continue
            header = try_clause_header(stripped)
            if header:
                nested_id, nested_title = header
                if nested_id.startswith(clause.clause_id + ".") or clause_level(nested_id) > clause_level(
                    clause.clause_id
                ):
                    pending_title = None
                    ensure(nested_id, nested_title, _sort_key_safe(nested_id))
                    current = nested_id
                    continue
            if not JUNK_LINE.search(stripped):
                # Short Hebrew/English line may be a title preceding an id-only marker.
                if len(stripped) <= 80 and not re.search(r"[.!?…:]$", stripped):
                    pending_title = stripped
                bodies.setdefault(current, []).append(stripped)

    from app.iso.parser import _sort_key

    result: list[ParsedClause] = []
    for cid in order:
        body = repair_broken_words(_clean_body("\n\n".join(bodies.get(cid, []))))
        result.append(ParsedClause(cid, by_id[cid].title, body, _sort_key(cid)))
    result.sort(key=lambda c: c.sort_order)
    return result


def _sort_key_safe(clause_id: str) -> int:
    from app.iso.parser import _sort_key

    return _sort_key(clause_id)


def ensure_parent_clauses(clauses: list[ParsedClause]) -> list[ParsedClause]:
    """Ensure parent ids exist when children were extracted (5.1 for 5.1.1)."""
    from app.iso.parser import _sort_key

    by_id = {c.clause_id: c for c in clauses}
    for cid in list(by_id):
        parts = cid.split(".")
        for i in range(1, len(parts)):
            parent = ".".join(parts[:i])
            if parent not in by_id:
                by_id[parent] = ParsedClause(parent, f"Clause {parent}", "", _sort_key(parent))
    return sorted(by_id.values(), key=lambda c: (c.sort_order, c.clause_id))


def parse_iso_document_text(raw: str, *, roll_up: bool = False) -> list[ParsedClause]:
    lines = reflow_continuation_lines(raw)
    # Always parse flat first, promote nested headings, then optionally roll up.
    clauses = parse_hierarchical_lines(lines, roll_up=False)
    if not clauses:
        raise ValueError(
            "No ISO clause headings found in document text. "
            "Ensure lines like '4.1 Understanding the organization' exist."
        )
    promoted = ensure_parent_clauses(promote_nested_clauses(clauses))
    # Keep parents even when title was missing in the source (common for HE/RTL).
    kept: list[ParsedClause] = []
    for clause in promoted:
        title = clause.title.strip() or f"Clause {clause.clause_id}"
        kept.append(ParsedClause(clause.clause_id, title, clause.body, clause.sort_order))
    if roll_up:
        kept = roll_up_empty_parents(kept)
    return kept


def _score_clause_set(clauses: list[ParsedClause]) -> tuple[int, int, int, int]:
    """Higher is better: intro present, level-3 count, core count, body chars."""
    core = []
    for c in clauses:
        top = c.clause_id.split(".", 1)[0]
        if top.isdigit() and 0 <= int(top) <= 10:
            core.append(c)
    level3 = sum(1 for c in core if c.clause_id.count(".") >= 2)
    body = sum(len((c.body or "").strip()) for c in core)
    intro = 1 if any(c.clause_id == "0" or c.clause_id.startswith("0.") for c in core) else 0
    return intro, level3, len(core), body


def parse_document_bytes(content: bytes, *, filename: str) -> list[ParsedClause]:
    """Single entry for PDF/DOC/DOCX uploads (regex-based, no LLM)."""
    from app.iso.document_extract import extract_text

    lower = filename.lower()
    candidates: list[list[ParsedClause]] = []

    # Text path includes Hebrew/RTL dotted-id normalization and level-3 promotion.
    try:
        text = extract_text(content, filename=filename)
        text = preprocess_extracted_text(text)
        candidates.append(parse_iso_document_text(text, roll_up=False))
    except Exception:
        pass

    if lower.endswith(".docx"):
        try:
            candidates.append(parse_docx_faithful(content))
        except Exception:
            pass
        try:
            candidates.append(parse_docx_by_structure(content))
        except Exception:
            pass
    elif lower.endswith((".pdf", ".doc")):
        if not candidates:
            raise ValueError("No ISO clause headings found in document text.")
    else:
        raise ValueError("Use parse_upload for .json/.md/.txt")

    candidates = [c for c in candidates if c]
    if not candidates:
        raise ValueError("No ISO clause headings found in document")
    return max(candidates, key=_score_clause_set)
