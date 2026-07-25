"""Tests for ISO parser and upload import."""
from __future__ import annotations

import json
import os

import pytest

os.environ["USE_SQLITE"] = "1"
os.environ["SQLITE_PATH"] = "/tmp/iso-parser-test.db"
os.environ["AUTH_DEV_MODE"] = "1"

from app.db import ensure_schema  # noqa: E402
from app.iso.parser import parse_upload  # noqa: E402
from app.iso.import_service import import_iso_upload  # noqa: E402
from app.db import get_conn  # noqa: E402
from app.main import app  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

client = TestClient(app)


@pytest.fixture(autouse=True)
def _schema():
    if os.path.exists(os.environ["SQLITE_PATH"]):
        os.remove(os.environ["SQLITE_PATH"])
    ensure_schema()


def test_sort_key_fits_postgres_integer():
    from app.iso.parser import _sort_key

    assert _sort_key("4.1") == 401000000
    assert _sort_key("10.1") == 1001000000
    assert _sort_key("4.1.2") == 401020000
    assert _sort_key("4") < _sort_key("4.1") < _sort_key("4.2") < _sort_key("5")
    assert _sort_key("4.1") < 2_147_483_647


def test_parse_markdown_clauses():
    raw = b"""4.1 Understanding the organization

The organization shall determine external and internal issues.

4.2 Needs of interested parties

The organization shall determine interested parties.
"""
    std, clauses = parse_upload(raw, filename="iso.md", standard="ISO9001", language="en")
    assert std == "ISO9001"
    by_id = {c.clause_id: c for c in clauses}
    # Parent "4" is synthesized so hierarchy / UI tree stays complete.
    assert set(by_id) == {"4", "4.1", "4.2"}
    assert "external and internal issues" in by_id["4.1"].body


def test_import_replaces_language_rows():
    content = json.dumps(
        [
            {"clause_id": "4.1", "title": "A", "body": "Body A", "sort_order": 401},
            {"clause_id": "4.2", "title": "B", "body": "Body B", "sort_order": 402},
        ]
    ).encode()
    with get_conn() as conn:
        import_iso_upload(
            conn,
            content=content,
            filename="iso.json",
            standard="ISO9001",
            language="en",
            edition="2015",
            admin_id=None,
        )
        conn.commit()
    token = client.post("/auth/dev-login", json={"email": "yaakovpreiger@gmail.com"}).json()["access_token"]
    r = client.get(
        "/v1/iso/clauses?standard=ISO9001&language=en",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r.status_code == 200
    assert len(r.json()["clauses"]) == 2


def test_import_bilingual_json():
    content = json.dumps(
        {
            "standard": "ISO9001",
            "clauses": [
                {
                    "clause_id": "4.1",
                    "sort_order": 401,
                    "en": {"title": "Understanding", "body": "English body"},
                    "he": {"title": "הבנת הארגון", "body": "גוף בעברית"},
                },
            ],
        }
    ).encode()
    with get_conn() as conn:
        result = import_iso_upload(
            conn,
            content=content,
            filename="iso-bilingual.json",
            standard="ISO9001",
            language="both",
            edition="2015",
            admin_id=None,
        )
        conn.commit()
    assert result["language"] == "both"
    assert result["clauses_en"] == 1
    assert result["clauses_he"] == 1
    token = client.post("/auth/dev-login", json={"email": "yaakovpreiger@gmail.com"}).json()["access_token"]
    en = client.get(
        "/v1/iso/clauses?standard=ISO9001&language=en",
        headers={"Authorization": f"Bearer {token}"},
    ).json()["clauses"]
    he = client.get(
        "/v1/iso/clauses?standard=ISO9001&language=he",
        headers={"Authorization": f"Bearer {token}"},
    ).json()["clauses"]
    assert any(c["clause_id"] == "4.1" and "English body" in c["text"] for c in en)
    assert any(c["clause_id"] == "4.1" and "גוף בעברית" in c["text"] for c in he)


def test_parse_docx_clauses():
    from docx import Document

    buf = __import__("io").BytesIO()
    doc = Document()
    doc.add_paragraph("4.1 Understanding the organization")
    doc.add_paragraph("The organization shall determine external and internal issues.")
    doc.add_paragraph("4.2 Needs of interested parties")
    doc.add_paragraph("The organization shall determine interested parties.")
    doc.save(buf)
    std, clauses = parse_upload(buf.getvalue(), filename="iso.docx", standard="ISO9001", language="en")
    assert std == "ISO9001"
    by_id = {c.clause_id: c for c in clauses}
    assert set(by_id) == {"4", "4.1", "4.2"}
    assert by_id["4.1"].title.startswith("Understanding")


def test_parse_pdf_clauses():
    from unittest.mock import patch

    sample = """4.1 Understanding the organization

The organization shall determine external and internal issues.

4.2 Needs of interested parties

The organization shall determine interested parties.
"""
    with patch("app.iso.document_extract.extract_text_from_pdf", return_value=sample):
        std, clauses = parse_upload(b"%PDF-1.4", filename="iso.pdf", standard="ISO9001", language="en")
    by_id = {c.clause_id: c for c in clauses}
    assert set(by_id) == {"4", "4.1", "4.2"}
    assert by_id["4.1"].clause_id == "4.1"


def test_dedupe_duplicate_clause_ids_in_upload():
    content = json.dumps(
        [
            {"clause_id": "1", "title": "First", "body": "A", "sort_order": 100},
            {"clause_id": "1", "title": "Second", "body": "B", "sort_order": 100},
            {"clause_id": "4.1", "title": "C", "body": "C body", "sort_order": 401},
        ]
    ).encode()
    with get_conn() as conn:
        result = import_iso_upload(
            conn,
            content=content,
            filename="iso.json",
            standard="ISO9001",
            language="en",
            edition="2015",
            admin_id=None,
            replace_previous=True,
        )
        conn.commit()
    assert result["clauses_imported"] == 2
    assert len(result["warnings"]) == 1
    token = client.post("/auth/dev-login", json={"email": "yaakovpreiger@gmail.com"}).json()["access_token"]
    clauses = client.get(
        "/v1/iso/clauses?standard=ISO9001&language=en",
        headers={"Authorization": f"Bearer {token}"},
    ).json()["clauses"]
    c1 = next(c for c in clauses if c["clause_id"] == "1")
    # First occurrence wins (avoids annex correspondence overwriting titles/bodies).
    assert "A" in c1["text"]
    assert "B" not in c1["text"]
    assert c1["title"] == "First"


def test_filters_pdf_noise():
    raw = (
        "000 Quality management systems\n\n"
        "Fundamentals junk\n\n"
        "4.1 Understanding the organization\n\n"
        "The organization shall determine issues.\n\n"
        "(c) ISO 2015 All rights reserved\n\n"
        "4.2 Needs of interested parties\n\n"
        "The organization shall determine interested parties.\n"
    ).encode()
    std, clauses = parse_upload(raw, filename="iso.txt", standard="ISO9001", language="en")
    assert std == "ISO9001"
    by_id = {c.clause_id: c for c in clauses}
    assert "0" not in by_id
    assert set(by_id) == {"4", "4.1", "4.2"}
    assert "All rights reserved" not in by_id["4.1"].body


def test_hebrew_view_uses_hebrew_rows():
    content = json.dumps(
        [
            {"clause_id": "4.1", "title": "HE title", "body": "גוף בעברית", "sort_order": 401},
        ]
    ).encode()
    with get_conn() as conn:
        import_iso_upload(
            conn,
            content=content,
            filename="iso-he.json",
            standard="ISO9001",
            language="he",
            edition="2015",
            admin_id=None,
            replace_previous=True,
        )
        conn.commit()
    token = client.post("/auth/dev-login", json={"email": "yaakovpreiger@gmail.com"}).json()["access_token"]
    r = client.get(
        "/v1/iso/clauses?standard=ISO9001&language=he",
        headers={"Authorization": f"Bearer {token}"},
    )
    c41 = next(c for c in r.json()["clauses"] if c["clause_id"] == "4.1")
    assert "גוף בעברית" in c41["text"]
    assert c41.get("fallback") is not True


def test_hebrew_fallback_to_english():
    token = client.post("/auth/dev-login", json={"email": "yaakovpreiger@gmail.com"}).json()["access_token"]
    r = client.get(
        "/v1/iso/clauses?standard=ISO9001&language=he",
        headers={"Authorization": f"Bearer {token}"},
    )
    clauses = r.json()["clauses"]
    assert len(clauses) >= 1
    # Seeded data has Hebrew — ensure API returns Hebrew text for 4.1
    c41 = next(c for c in clauses if c["clause_id"] == "4.1")
    assert "ארגון" in c41["text"] or c41.get("fallback")


def test_seed_data_sort_order_scale():
    """Verify seed data uses new 9-digit scale."""
    from pathlib import Path
    from app.iso.parser import _sort_key

    seed_path = Path(__file__).parent.parent / "app/data/iso_clauses_seed.json"
    data = json.load(seed_path.open())

    for item in data:
        expected = _sort_key(item["clause_id"])
        actual = item["sort_order"]
        assert actual == expected, \
            f"Clause {item['clause_id']}: sort_order {actual} != expected {expected}"


def test_clause_ordering_consistency():
    """Verify clause ordering matches natural hierarchy."""
    from app.iso.parser import _sort_key

    content = json.dumps(
        [
            {"clause_id": "4", "title": "Context", "body": "4 body", "sort_order": _sort_key("4")},
            {"clause_id": "4.1", "title": "Understanding", "body": "4.1 body", "sort_order": _sort_key("4.1")},
            {"clause_id": "4.1.1", "title": "Sub", "body": "4.1.1 body", "sort_order": _sort_key("4.1.1")},
            {"clause_id": "4.2", "title": "Needs", "body": "4.2 body", "sort_order": _sort_key("4.2")},
            {"clause_id": "5", "title": "Leadership", "body": "5 body", "sort_order": _sort_key("5")},
        ]
    ).encode()

    with get_conn() as conn:
        import_iso_upload(
            conn,
            content=content,
            filename="test-order.json",
            standard="TESTORDER",
            language="en",
            edition="2025",
            admin_id=None,
            replace_previous=True,
        )
        conn.commit()

    token = client.post("/auth/dev-login", json={"email": "yaakovpreiger@gmail.com"}).json()["access_token"]
    r = client.get(
        "/v1/iso/clauses?standard=TESTORDER&language=en",
        headers={"Authorization": f"Bearer {token}"},
    )

    clause_ids = [c["clause_id"] for c in r.json()["clauses"]]
    expected_order = ["4", "4.1", "4.1.1", "4.2", "5"]
    assert clause_ids == expected_order, f"Expected {expected_order}, got {clause_ids}"
