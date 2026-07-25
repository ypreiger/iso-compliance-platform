"""Document text extraction from PDF, DOC, DOCX, and Excel.

VENDORED MODULE: This code is duplicated in apps/iso-api/app/iso/document_extract.py
for service isolation. Bug fixes must be applied to BOTH locations.

This layer is purely mechanical — it extracts raw text without interpretation.
The LLM extractor (extractor.py) handles semantic structuring.
"""
from __future__ import annotations

import logging
import os
import re
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

log = logging.getLogger(__name__)

HEBREW_CHAR_RE = re.compile(r"[\u0590-\u05FF]")
CLAUSE_ID_TOKEN_RE = re.compile(r"^\d{1,2}(?:\.\d{1,2}){0,4}[\.\)]?$")
HEBREW_PREFIX_LETTERS = {"ו", "ב", "כ", "ל", "מ", "ש", "ה"}

_SPARSE_AFTER_HEADING = re.compile(
    r"(?m)^(\d{1,2}(?:\.\d{1,2}){1,4})\b[^\n]{0,160}\n"
    r"(?:shall\b|must\b|maintained\b|c\)|procedures shall\b|other relevant\b|"
    r"documentation,|authorizing its use)",
    re.I,
)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def _contains_hebrew(text: str) -> bool:
    return bool(HEBREW_CHAR_RE.search(text))


def _hebrew_ratio(text: str) -> float:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return 0.0
    hebrew = sum(1 for c in letters if HEBREW_CHAR_RE.match(c))
    return hebrew / len(letters)


def _is_single_hebrew_letter(token: str) -> bool:
    return len(token) == 1 and bool(HEBREW_CHAR_RE.match(token))


def _starts_with_hebrew(token: str) -> bool:
    for char in token:
        if char.isalpha():
            return bool(HEBREW_CHAR_RE.match(char))
    return False


def _fix_hebrew_spacing(text: str) -> str:
    """Conservative fix for Hebrew token fragmentation."""
    if not text or not _contains_hebrew(text):
        return text

    fixed_lines: list[str] = []
    for line in text.split("\n"):
        if not _contains_hebrew(line):
            fixed_lines.append(line)
            continue
        words = line.split()
        rebuilt: list[str] = []
        i = 0
        while i < len(words):
            word = words[i]
            if _is_single_hebrew_letter(word):
                run = [word]
                j = i + 1
                while j < len(words) and _is_single_hebrew_letter(words[j]):
                    run.append(words[j])
                    j += 1
                if len(run) >= 2:
                    rebuilt.append("".join(run))
                    i = j
                    continue
                if (
                    j < len(words)
                    and word in HEBREW_PREFIX_LETTERS
                    and _starts_with_hebrew(words[j])
                ):
                    rebuilt.append(word + words[j])
                    i = j + 1
                    continue
            rebuilt.append(word)
            i += 1
        fixed_lines.append(" ".join(rebuilt))
    return "\n".join(fixed_lines)


def _is_clause_id_token(token: str) -> bool:
    return bool(CLAUSE_ID_TOKEN_RE.match(token))


def _reconstruct_hebrew_line(entries: list[tuple[float, str]]) -> str:
    if not entries:
        return ""
    tokens = [word for _, word in sorted(entries, key=lambda item: item[0])]
    sample = " ".join(tokens)
    if _hebrew_ratio(sample) < 0.4:
        return " ".join(tokens)
    prefix_len = 0
    for tok in tokens:
        if _is_clause_id_token(tok):
            prefix_len += 1
            continue
        break
    ordered = tokens[:prefix_len] + list(reversed(tokens[prefix_len:]))
    return _fix_hebrew_spacing(" ".join(ordered))


def _extract_hebrew_page_text(page) -> str:
    words_list = page.get_text("words", sort=False)
    if not words_list:
        return page.get_text("text", sort=True)

    words_by_block: dict[int, dict[int, list[tuple[float, float, str]]]] = {}
    for item in words_list:
        if len(item) < 8:
            continue
        x0, y0, _x1, _y1, word, block_no, line_no, _word_no = item[:8]
        word_text = str(word).strip()
        if not word_text:
            continue
        words_by_block.setdefault(int(block_no), {}).setdefault(int(line_no), []).append(
            (float(x0), float(y0), word_text)
        )

    rendered_blocks: list[str] = []
    seen_blocks: set[int] = set()
    for block in page.get_text("blocks", sort=True):
        if len(block) < 7 or int(block[6]) != 0:
            continue
        block_no = int(block[5])
        seen_blocks.add(block_no)
        line_map = words_by_block.get(block_no)
        if not line_map:
            fallback = str(block[4]).strip()
            if fallback:
                rendered_blocks.append(fallback)
            continue
        rendered_lines: list[tuple[float, int, str]] = []
        for line_no, words in line_map.items():
            line = _reconstruct_hebrew_line([(x, text) for x, _y, text in words]).strip()
            if line:
                y0 = min(y for _x, y, _text in words)
                rendered_lines.append((y0, line_no, line))
        rendered_lines.sort(key=lambda item: (item[0], item[1]))
        if rendered_lines:
            rendered_blocks.append("\n".join(line for _y, _ln, line in rendered_lines))

    for block_no, line_map in words_by_block.items():
        if block_no in seen_blocks:
            continue
        rendered_lines: list[tuple[float, int, str]] = []
        for line_no, words in line_map.items():
            line = _reconstruct_hebrew_line([(x, text) for x, _y, text in words]).strip()
            if line:
                y0 = min(y for _x, y, _text in words)
                rendered_lines.append((y0, line_no, line))
        rendered_lines.sort(key=lambda item: (item[0], item[1]))
        if rendered_lines:
            rendered_blocks.append("\n".join(line for _y, _ln, line in rendered_lines))

    return "\n\n".join(rendered_blocks)


# ── PDF OCR (sparse text-layer recovery) ───────────────────────────────────

def _ocr_enabled() -> bool:
    return os.getenv("ISO_PDF_OCR", "auto").strip().lower() not in {"0", "false", "off", "no"}


def page_text_looks_sparse(text: str) -> bool:
    raw = text or ""
    compact = re.sub(r"\s+", "", raw)
    has_clause = bool(re.search(r"(?m)^\d{1,2}(?:\.\d{1,2}){1,4}\b", raw))
    if has_clause and len(compact) < 1200:
        return True
    return bool(_SPARSE_AFTER_HEADING.search(raw))


def _ocr_pdf_page(page) -> str:
    try:
        tp = page.get_textpage_ocr(language="eng", dpi=250, full=True)
        text = (page.get_text("text", textpage=tp) or "").strip()
        if text and len(re.sub(r"\s+", "", text)) > 200:
            return text
    except Exception as exc:
        log.debug("tesseract OCR failed: %s", exc)
    try:
        import numpy as np
        import easyocr  # type: ignore

        pix = page.get_pixmap(dpi=250)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        if pix.n >= 4:
            img = img[:, :, :3]
        reader = getattr(_ocr_pdf_page, "_reader", None)
        if reader is None:
            reader = easyocr.Reader(["en"], gpu=False, verbose=False)
            _ocr_pdf_page._reader = reader  # type: ignore[attr-defined]
        return "\n".join(reader.readtext(img, detail=0, paragraph=True)).strip()
    except Exception as exc:
        log.warning("PDF OCR unavailable: %s", exc)
        return ""


# ── PDF ────────────────────────────────────────────────────────────────────

def extract_pdf(data: bytes) -> str:
    """Extract text from PDF in reading order using PyMuPDF (fitz).

    For Hebrew PDFs, uses word-based extraction for better quality.
    For sparse English Print-to-PDF layers, OCR recovers missing clause bodies.
    """
    is_hebrew = False
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=data, filetype="pdf")

        # Check if Hebrew content (sample first 2 pages)
        for page_num in range(min(2, len(doc))):
            sample = doc[page_num].get_text("text")[:400]
            if sum(1 for c in sample if HEBREW_CHAR_RE.match(c)) > 15:
                is_hebrew = True
                break

        pages: list[str] = []

        if is_hebrew:
            # Preserve block/line structure and reconstruct RTL lines.
            for page in doc:
                text = _extract_hebrew_page_text(page)
                if text.strip():
                    pages.append(text)
        else:
            for page in doc:
                text = page.get_text("text", sort=True) or ""
                if not text.strip():
                    blocks = page.get_text("blocks", sort=True)
                    text = "\n".join(str(b[4]) for b in blocks if b[6] == 0)
                if _ocr_enabled() and page_text_looks_sparse(text):
                    ocr_text = _ocr_pdf_page(page)
                    if ocr_text and len(re.sub(r"\s+", "", ocr_text)) > len(re.sub(r"\s+", "", text)) * 1.2:
                        text = ocr_text
                if text.strip():
                    pages.append(text)

        doc.close()
        text = "\n".join(pages)
    except ImportError:
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)

    if not text.strip():
        raise ValueError("PDF contains no extractable text (scanned PDF not supported)")
    if is_hebrew or _contains_hebrew(text):
        text = _fix_hebrew_spacing(text)
    return _normalize(text)


# ── DOCX ───────────────────────────────────────────────────────────────────

def extract_docx(data: bytes) -> str:
    """Extract text from DOCX; prefix headings with ## for downstream detection."""
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(BytesIO(data))
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            para = Paragraph(child, doc)
            text = re.sub(r"\s+", " ", para.text).strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            lines.append(f"## {text}" if "heading" in style else text)
        elif isinstance(child, CT_Tbl):
            table = Table(child, doc)
            for row in table.rows:
                cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
    return _normalize("\n".join(lines))


# ── DOC (legacy) ───────────────────────────────────────────────────────────

def extract_doc(data: bytes) -> str:
    """Extract text from legacy .doc via antiword binary."""
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "upload.doc"
        p.write_bytes(data)
        env = {"ANTIWORDHOME": "/usr/share/antiword", **os.environ}
        try:
            res = subprocess.run(
                ["antiword", "-m", "UTF-8.txt", "-w", "0", str(p)],
                check=True, capture_output=True, timeout=120, env=env,
            )
        except FileNotFoundError as e:
            raise ValueError("antiword not installed in container") from e
        except subprocess.CalledProcessError:
            # Retry without explicit mapping when map files are missing.
            try:
                res = subprocess.run(
                    ["antiword", "-w", "0", str(p)],
                    check=True, capture_output=True, timeout=120, env=env,
                )
            except subprocess.CalledProcessError as e:
                raise ValueError(f".doc extraction failed: {(e.stderr or b'').decode()[:300]}") from e
        text = res.stdout.decode("utf-8", errors="replace")
        if not text.strip():
            raise ValueError(".doc file has no extractable text")
        return _normalize(text)


# ── Excel ──────────────────────────────────────────────────────────────────

def extract_excel(data: bytes) -> dict:
    """Parse Excel file; return sheet data as list of dicts (headers as keys)."""
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheets: list[dict] = []
    for ws in wb.worksheets:
        rows_raw = list(ws.iter_rows(values_only=True))
        if not rows_raw:
            continue
        headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows_raw[0])]
        rows: list[dict] = []
        for row in rows_raw[1:]:
            if all(v is None for v in row):
                continue
            rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        sheets.append({"name": ws.title, "headers": headers, "rows": rows})
    wb.close()
    return {"sheets": sheets}


# ── unified ────────────────────────────────────────────────────────────────

def extract_text(data: bytes, filename: str) -> str:
    """Extract plain text from any supported document type."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_pdf(data)
    if lower.endswith(".docx"):
        return extract_docx(data)
    if lower.endswith(".doc"):
        return extract_doc(data)
    raise ValueError(f"Unsupported file type for text extraction: {filename}. Use .pdf, .doc, or .docx")
