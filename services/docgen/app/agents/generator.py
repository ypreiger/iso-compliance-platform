"""Document generation: Excel, DOCX, PDF output.

All generators are synchronous (heavy I/O) and return raw bytes.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any


# ── Excel ──────────────────────────────────────────────────────────────────

def generate_excel_compliance_report(data: dict) -> bytes:
    """Generate an Excel compliance report from project data.

    data keys:
        project_name, standard, edition, locale,
        findings: [{id, text, severity}]
        mappings: [{finding_id, clause_id, clause_title, relevance_pct, severity}]
        coverage: [{clause_id, clause_title, status, reason}]
    """
    import xlsxwriter

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})

    is_he = data.get("locale", "en") == "he"

    # styles
    hdr = wb.add_format({"bold": True, "bg_color": "#003087", "font_color": "#FFFFFF", "border": 1})
    cell = wb.add_format({"border": 1, "text_wrap": True})
    title_fmt = wb.add_format({"bold": True, "font_size": 14})
    ok_fmt = wb.add_format({"border": 1, "bg_color": "#d4edda"})
    warn_fmt = wb.add_format({"border": 1, "bg_color": "#fff3cd"})
    fail_fmt = wb.add_format({"border": 1, "bg_color": "#f8d7da"})

    # ── Sheet 1: Summary ───────────────────────────────────────────────────
    ws_sum = wb.add_worksheet("Summary" if not is_he else "סיכום")
    ws_sum.write(0, 0, data.get("project_name", "Compliance Report"), title_fmt)
    ws_sum.write(1, 0, f"Standard: {data.get('standard', '')} | Edition: {data.get('edition', '')} | Date: {date.today()}")
    ws_sum.write(3, 0, "Findings", hdr)
    ws_sum.write(3, 1, str(len(data.get("findings", []))), cell)
    ws_sum.write(4, 0, "Mapped clauses", hdr)
    ws_sum.write(4, 1, str(len({m["clause_id"] for m in data.get("mappings", [])})), cell)
    ws_sum.write(5, 0, "Covered clauses", hdr)
    covered = [c for c in data.get("coverage", []) if c.get("status") == "covered"]
    ws_sum.write(5, 1, str(len(covered)), cell)
    ws_sum.set_column(0, 0, 25)
    ws_sum.set_column(1, 1, 40)

    # ── Sheet 2: Findings ──────────────────────────────────────────────────
    ws_f = wb.add_worksheet("Findings" if not is_he else "ממצאים")
    hdrs_f = ["ID", "Finding", "Severity"] if not is_he else ["מזהה", "ממצא", "חומרה"]
    for col, h in enumerate(hdrs_f):
        ws_f.write(0, col, h, hdr)
    for i, f in enumerate(data.get("findings", []), start=1):
        ws_f.write(i, 0, f.get("id", ""), cell)
        ws_f.write(i, 1, f.get("text", ""), cell)
        sev = f.get("severity", "minor")
        fmt = fail_fmt if sev == "major" else warn_fmt if sev == "moderate" else cell
        ws_f.write(i, 2, sev, fmt)
    ws_f.set_column(0, 0, 10)
    ws_f.set_column(1, 1, 60)
    ws_f.set_column(2, 2, 12)

    # ── Sheet 3: Clause Mappings ───────────────────────────────────────────
    ws_m = wb.add_worksheet("Mappings" if not is_he else "מיפויים")
    hdrs_m = ["Finding ID", "Clause ID", "Clause Title", "Relevance %", "Severity"] \
        if not is_he else ["מזהה ממצא", "מזהה סעיף", "כותרת סעיף", "רלוונטיות %", "חומרה"]
    for col, h in enumerate(hdrs_m):
        ws_m.write(0, col, h, hdr)
    for i, m in enumerate(data.get("mappings", []), start=1):
        ws_m.write(i, 0, m.get("finding_id", ""), cell)
        ws_m.write(i, 1, m.get("clause_id", ""), cell)
        ws_m.write(i, 2, m.get("clause_title", ""), cell)
        ws_m.write(i, 3, m.get("relevance_pct", 0), cell)
        sev = m.get("severity", "minor")
        fmt = fail_fmt if sev == "major" else warn_fmt if sev == "moderate" else cell
        ws_m.write(i, 4, sev, fmt)
    ws_m.set_column(0, 0, 12)
    ws_m.set_column(1, 1, 10)
    ws_m.set_column(2, 2, 45)
    ws_m.set_column(3, 3, 14)
    ws_m.set_column(4, 4, 12)

    # ── Sheet 4: Coverage Matrix ───────────────────────────────────────────
    ws_c = wb.add_worksheet("Coverage" if not is_he else "כיסוי")
    hdrs_c = ["Clause ID", "Clause Title", "Status", "Notes"] \
        if not is_he else ["מזהה סעיף", "כותרת סעיף", "סטטוס", "הערות"]
    for col, h in enumerate(hdrs_c):
        ws_c.write(0, col, h, hdr)
    status_fmt = {"covered": ok_fmt, "partial": warn_fmt, "not_covered": fail_fmt}
    for i, c in enumerate(data.get("coverage", []), start=1):
        ws_c.write(i, 0, c.get("clause_id", ""), cell)
        ws_c.write(i, 1, c.get("clause_title", ""), cell)
        st = c.get("status", "not_covered")
        ws_c.write(i, 2, st, status_fmt.get(st, cell))
        ws_c.write(i, 3, c.get("reason", ""), cell)
    ws_c.set_column(0, 0, 10)
    ws_c.set_column(1, 1, 45)
    ws_c.set_column(2, 2, 14)
    ws_c.set_column(3, 3, 50)

    wb.close()
    buf.seek(0)
    return buf.read()


# ── DOCX ───────────────────────────────────────────────────────────────────

def generate_docx_report(data: dict) -> bytes:
    """Generate a Word (.docx) compliance report."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    is_he = data.get("locale", "en") == "he"

    # Title
    title = doc.add_heading(data.get("project_name", "Compliance Report"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT if is_he else WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph(
        f"Standard: {data.get('standard', '')} | Edition: {data.get('edition', '')} | Date: {date.today()}"
    )

    # Findings
    doc.add_heading("Findings" if not is_he else "ממצאים", level=1)
    if data.get("findings"):
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        hdr_row = tbl.rows[0].cells
        for i, h in enumerate(["ID", "Description", "Severity"] if not is_he else ["מזהה", "תיאור", "חומרה"]):
            hdr_row[i].text = h
        for f in data.get("findings", []):
            row = tbl.add_row().cells
            row[0].text = str(f.get("id", ""))
            row[1].text = str(f.get("text", ""))
            row[2].text = str(f.get("severity", ""))
    else:
        doc.add_paragraph("No findings." if not is_he else "אין ממצאים.")

    # Clause mappings
    doc.add_heading("Clause Mappings" if not is_he else "מיפויים לסעיפים", level=1)
    if data.get("mappings"):
        tbl2 = doc.add_table(rows=1, cols=4)
        tbl2.style = "Table Grid"
        hdr2 = tbl2.rows[0].cells
        for i, h in enumerate(["Finding", "Clause", "Title", "Relevance"] if not is_he else ["ממצא", "סעיף", "כותרת", "רלוונטיות"]):
            hdr2[i].text = h
        for m in data.get("mappings", []):
            row = tbl2.add_row().cells
            row[0].text = str(m.get("finding_id", ""))
            row[1].text = str(m.get("clause_id", ""))
            row[2].text = str(m.get("clause_title", ""))
            row[3].text = f"{m.get('relevance_pct', 0)}%"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── PDF ────────────────────────────────────────────────────────────────────

def generate_pdf_report(data: dict) -> bytes:
    """Generate a PDF compliance report using ReportLab."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )

    buf = io.BytesIO()
    is_he = data.get("locale", "en") == "he"
    doc_rl = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm,
                                topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    story: list[Any] = []

    # Title
    story.append(Paragraph(f"<b>{data.get('project_name', 'Compliance Report')}</b>",
                            styles["Title"]))
    story.append(Paragraph(
        f"Standard: {data.get('standard', '')} | Edition: {data.get('edition', '')} | {date.today()}",
        styles["Normal"],
    ))
    story.append(Spacer(1, 0.5*cm))

    # Summary table
    sum_data = [
        ["Metric", "Value"],
        ["Total findings", str(len(data.get("findings", [])))],
        ["Mapped clauses", str(len({m["clause_id"] for m in data.get("mappings", [])}))],
        ["Covered clauses", str(len([c for c in data.get("coverage", []) if c.get("status") == "covered"]))],
    ]
    t = Table(sum_data, colWidths=[8*cm, 5*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003087")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    # Findings
    story.append(Paragraph("<b>Findings</b>", styles["Heading2"]))
    if data.get("findings"):
        f_data = [["ID", "Description", "Severity"]]
        for f in data["findings"]:
            f_data.append([str(f.get("id", "")), str(f.get("text", ""))[:120], str(f.get("severity", ""))])
        ft = Table(f_data, colWidths=[2*cm, 12*cm, 3*cm])
        ft.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003087")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(ft)

    doc_rl.build(story)
    buf.seek(0)
    return buf.read()
